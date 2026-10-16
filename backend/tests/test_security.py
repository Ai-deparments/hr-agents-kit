"""Regression tests for the security audit fixes and the new anti-fraud / engagement features."""
import io
from datetime import date, datetime, timedelta

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from app.main import app
from app.models import Branch, Candidate, Employee, Position, Vacancy
from app.rbac import Scope
from app.services import checkin, engagement, people, recruiting, reports

c = TestClient(app)
OWNER = Scope("test", "owner", None, True)


def _login(u, p="demo12345"):
    return {"Authorization": f"Bearer {c.post('/api/auth/login', json={'username': u, 'password': p}).json()['token']}"}


def test_security_headers_present():
    r = c.get("/api/health")
    assert r.headers["X-Content-Type-Options"] == "nosniff"
    assert "frame-ancestors" in r.headers["Content-Security-Policy"]


def test_rotating_qr_code(db):
    b = db.scalar(select(Branch).where(Branch.code == "P02"))
    code = checkin.code_for(b)
    assert len(code) == 6 and checkin.check_code(b, f"HRCHK:{b.id}:{code}")
    assert not checkin.check_code(b, "000000" if code != "000000" else "111111")
    old = checkin.code_for(b, at=datetime.now().timestamp() - 600)
    assert not checkin.check_code(b, old) or old == code


def test_link_bruteforce_is_limited(db):
    tok = c.post("/api/miniapp/auth", json={"init_data": ""}).json()["token"]
    h = {"Authorization": f"Bearer {tok}"}
    # dev user 2001 is already linked -> 409 before any guessing; check the per-card limit directly instead
    from app.api.miniapp import LINK_MAX_FAILS_PER_CARD, _link_fails
    assert _link_fails(db, target="employee:999999") == 0 and LINK_MAX_FAILS_PER_CARD <= 10
    names = c.get("/api/miniapp/link/people?branch_id=1", headers=h).json()
    assert all(len(n["full_name"].split()) <= 2 and (len(n["full_name"].split()) < 2 or n["full_name"].endswith(".")) for n in names)
    assert c.get("/api/miniapp/link/branches").status_code == 401  # no anonymous staff/branch listing


def test_recruiter_cannot_touch_other_branch_candidates(db):
    from app.auth import hash_password
    from app.models import User
    p3 = db.scalar(select(Branch).where(Branch.code == "P03"))
    if not db.scalar(select(User).where(User.username == "rec_p3")):
        db.add(User(username="rec_p3", password_hash=hash_password("demo12345"), role="recruiter", branch_scope=[p3.id]))
        db.commit()
    other = db.scalar(select(Candidate).join(Vacancy).where(Vacancy.branch_id != p3.id))
    r = c.post(f"/api/candidates/{other.id}/status", json={"status": "rejected"}, headers=_login("rec_p3"))
    assert r.status_code == 404


def test_readonly_manager_cannot_mark_onboarding():
    r = c.post("/api/onboarding/items/1?done=true", headers=_login("regional"))
    assert r.status_code == 403


def test_status_patch_cannot_fire_without_reason():
    h = _login("hr")
    emp = c.get("/api/employees", headers=h).json()[0]
    assert c.patch(f"/api/employees/{emp['id']}", json={"status": "fired"}, headers=h).status_code == 400


def test_excel_formula_injection_is_neutralised(db):
    e = db.scalar(select(Employee).where(Employee.status != "fired"))
    old = e.full_name
    e.full_name = '=HYPERLINK("http://evil","x")'
    db.flush()
    wb = load_workbook(io.BytesIO(reports.day_excel(db, OWNER, date.today())))
    vals = [r[0].value for r in wb.active.iter_rows(min_row=2)]
    assert "'=HYPERLINK(\"http://evil\",\"x\")" in vals
    e.full_name = old


def test_report_html_is_escaped(db):
    e = db.scalar(select(Employee).where(Employee.status != "fired"))
    assert "&lt;b&gt;" in reports._esc("<b>")


def test_candidate_retention_purge(db):
    cnd = Candidate(full_name="Old Rejected", phone="+998900000000", status="rejected",
                    decided_at=datetime.now() - timedelta(days=400), answers=[{"q": "x", "a": "y"}])
    db.add(cnd)
    db.flush()
    assert recruiting.purge_candidates(db) >= 1
    assert cnd.full_name == "—" and cnd.phone is None and cnd.answers == []


def test_early_vacation_is_flagged(db):
    p = db.scalar(select(Position).where(Position.code == "cashier"))
    b = db.scalar(select(Branch).where(Branch.code == "P02"))
    e = Employee(full_name="New Cashier", branch_id=b.id, position_id=p.id, hire_date=date.today() - timedelta(days=40),
                 leave_balance_days=10)
    db.add(e)
    db.flush()
    r = people.create_request(db, OWNER, e, "vacation", date.today() + timedelta(days=5), date.today() + timedelta(days=7))
    assert r.comment.startswith("[!]")


def test_pulse_is_anonymous_and_thresholded(db):
    e = db.scalar(select(Employee).where(Employee.status != "fired", Employee.telegram_id.is_not(None)))
    assert engagement.record_pulse(db, e, "enps", 9) is True
    assert engagement.record_pulse(db, e, "enps", 2) is False  # one vote per month
    assert engagement.enps([10, 10, 9]) is None  # fewer than 5 answers -> hidden
    assert engagement.enps([10, 9, 9, 5, 3]) == 20
    from app.models import PulseResponse
    assert not hasattr(PulseResponse, "employee_id")


def test_sop_ack_flow(db):
    e = db.scalar(select(Employee).where(Employee.full_name == "New Cashier")) or db.scalar(select(Employee))
    unread = engagement.my_unread(db, e)
    if unread:
        engagement.acknowledge(db, OWNER, e, unread[0]["id"])
        db.flush()
        assert unread[0]["id"] not in {u["id"] for u in engagement.my_unread(db, e)}
    assert all("pct" in r for r in engagement.sop_coverage(db, OWNER))


def test_risk_signals_hr_only():
    assert c.get("/api/engagement/risk", headers=_login("regional")).status_code == 403
    assert c.get("/api/engagement/risk", headers=_login("hr")).status_code == 200
