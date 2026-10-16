"""Deterministic reports (no LLM): same format, same time, unknown = '—'."""
from __future__ import annotations

import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from ..config import cfg
from ..rbac import Scope
from . import attendance as att
from . import people

STATUS_LABEL = {"present": "came", "late": "late", "on_shift": "on shift", "absent": "absent", "leave": "leave",
                "unmarked": "not marked", "off": "day off", "needs_review": "review", "conflict": "conflict"}


def _bar(part: int, total: int, width: int = 10) -> str:
    if not total:
        return "—"
    n = round(width * part / total)
    return "█" * n + "░" * (width - n)


def morning_text(db: Session, scope: Scope, d: date | None = None, now: datetime | None = None) -> str:
    d = d or date.today()
    s = att.day_summary(db, scope, d, now=now)
    brand = cfg("branding.short", "HR")
    lines = [f"🌅 <b>MORNING ATTENDANCE</b> · {brand} · {d.isoformat()}", "",
             f"Scheduled: <b>{s['scheduled']}</b> · Came: <b>{s['came']}</b> · Late: <b>{s['late']}</b> · Not marked: <b>{s['unmarked']}</b>"]
    problems = {k: v for k, v in s["by_branch"].items() if v.get("unmarked") or v.get("late") or not v.get("came")}
    if problems:
        lines += ["", "<b>Branches needing attention</b>"]
        for name, v in list(problems.items())[:15]:
            sched = v["total"] - v.get("off", 0)
            lines.append(f"{_bar(v.get('came', 0), sched)} {name}: {v.get('came', 0)}/{sched}, late {v.get('late', 0)}")
    late = [r for r in s["rows"] if r["status"] == "late"][:10]
    if late:
        lines += ["", "<b>Late</b>"] + [f"• {r['full_name']} ({r['branch']}) — {r['check_in'][11:]} (+{r['late_min']} min)" for r in late]
    lines += ["", "<i>not marked ≠ absent</i>"]
    return "\n".join(lines)


def evening_text(db: Session, scope: Scope, d: date | None = None, now: datetime | None = None) -> str:
    d = d or date.today()
    s = att.day_summary(db, scope, d, now=now)
    worked = [r["worked_hours"] for r in s["rows"] if r["worked_hours"]]
    avg = f"{sum(worked) / len(worked):.1f} h" if worked else "—"
    lines = [f"🌙 <b>DAY SUMMARY</b> · {d.isoformat()}", "",
             f"Came: <b>{s['came']}</b> of {s['scheduled']} · Late: <b>{s['late']}</b> · On shift (24h/overnight): <b>{s['on_shift']}</b>",
             f"Needs review: <b>{s['needs_review'] + s['conflict']}</b> · Avg worked: <b>{avg}</b>"]
    for name, v in list(s["by_branch"].items())[:15]:
        sched = v["total"] - v.get("off", 0)
        lines.append(f"{_bar(v.get('came', 0), sched)} {name}: {v.get('came', 0)}/{sched}")
    lines += ["", "<i>overnight shifts check out tomorrow — shown as 'on shift', not missing</i>"]
    return "\n".join(lines)


def day_excel(db: Session, scope: Scope, d: date) -> bytes:
    s = att.day_summary(db, scope, d)
    wb = Workbook()
    ws = wb.active
    ws.title = d.isoformat()
    ws.append(["Full name", "Branch", "Position", "Status", "Check-in", "Check-out", "Worked h", "Late min", "Source"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in s["rows"]:
        ws.append([r["full_name"], r["branch"], r["position"], STATUS_LABEL.get(r["status"], r["status"]),
                   r["check_in"] or "—", r["check_out"] or "—", r["worked_hours"] if r["worked_hours"] is not None else "—",
                   r["late_min"] or "", r["source"] or "—"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def tabel_excel(db: Session, scope: Scope, month: str, branch_id: int | None = None) -> bytes:
    t = att.month_tabel(db, scope, month, branch_id)
    wb = Workbook()
    ws = wb.active
    ws.title = f"tabel-{month}"
    ws.append(["Full name", "Branch", "Position"] + t["days"] + ["Worked h", "Norm h", "Days", "Late", "Review", "Unknown"])
    for c in ws[1]:
        c.font = Font(bold=True)
    fills = {"L": "FDE68A", "A": "FCA5A5", "V": "BFDBFE", "R": "F9A8D4", "!": "F87171", "?": "E5E7EB"}
    for row in t["rows"]:
        ws.append([row["full_name"], row["branch"], row["position"]] +
                  [(c["h"] if c["h"] else c["c"]) for c in row["cells"]] +
                  [row["worked_hours"], row["norm_hours"], row["worked_days"], row["late_days"], row["review_days"], row["unknown_days"]])
        for i, c in enumerate(row["cells"]):
            if c["c"] in fills:
                ws.cell(ws.max_row, 4 + i).fill = PatternFill("solid", fgColor=fills[c["c"]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def documents_digest(db: Session, scope: Scope, days: int = 30) -> str | None:
    docs = people.expiring_documents(db, scope, days)
    if not docs:
        return None
    lines = [f"📄 <b>Documents expiring in {days} days</b>"]
    for d in docs[:20]:
        when = "EXPIRED" if d["expired"] else f"{d['days_left']} d"
        lines.append(f"• {d['full_name']} ({d['branch']}) — {d['title']}: {when}")
    return "\n".join(lines)
